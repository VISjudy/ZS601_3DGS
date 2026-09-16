#!/usr/bin/env python3
"""Build the evidence-backed ZS601 experiment record.

The generator treats the JSON registry as the only structured source for run
status and measured values. Null or absent values are rendered as 待核验; it
never invents metrics from plans or directory names.
"""

from __future__ import annotations

import argparse
import json
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterable

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_REGISTRY = ROOT / "experiments" / "experiment_registry.json"
DEFAULT_OUTPUT = ROOT / "reports" / "ZS601_3DGS_实验总记录.docx"


def verified(value: Any) -> str:
    if value is None or value == "" or value == [] or value == {}:
        return "待核验"
    if isinstance(value, bool):
        return "是" if value else "否"
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False, sort_keys=True)
    return str(value)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cell_borders(cell, color: str = "D9D9D9") -> None:
    tc_pr=cell._tc.get_or_add_tcPr()
    borders=tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders=OxmlElement("w:tcBorders"); tc_pr.append(borders)
    for edge in ("top","left","bottom","right","insideH","insideV"):
        tag="w:"+edge
        node=borders.find(qn(tag))
        if node is None:
            node=OxmlElement(tag); borders.append(node)
        node.set(qn("w:val"),"single")
        node.set(qn("w:sz"),"4")
        node.set(qn("w:color"),color)


def evidenced(exp: dict[str, Any], key: str) -> Any:
    return exp.get(key) if exp.get("evidence") else None


def set_cell_text(cell, value: Any, bold: bool = False) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(verified(value))
    run.bold = bold
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(9.5)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(doc: Document, headers: list[str], rows: Iterable[Iterable[Any]]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    header = table.rows[0]
    repeat_table_header(header)
    for cell, value in zip(header.cells, headers):
        set_cell_text(cell, value, bold=True)
        set_cell_shading(cell, "D9EAF7")
    for row_index, values in enumerate(rows):
        row = table.add_row()
        for cell, value in zip(row.cells, values):
            set_cell_text(cell, value)
            if row_index % 2:
                set_cell_shading(cell, "F7F7F7")
    for row in table.rows:
        for cell in row.cells:
            set_cell_borders(cell)
            tc_pr=cell._tc.get_or_add_tcPr()
            tc_mar=tc_pr.first_child_found_in("w:tcMar")
            if tc_mar is None:
                tc_mar=OxmlElement("w:tcMar"); tc_pr.append(tc_mar)
            for side in ("top","left","bottom","right"):
                node=tc_mar.find(qn("w:"+side))
                if node is None:
                    node=OxmlElement("w:"+side); tc_mar.append(node)
                node.set(qn("w:w"),"90"); node.set(qn("w:type"),"dxa")
    doc.add_paragraph()


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    heading = doc.add_heading(text, level=level)
    heading.paragraph_format.keep_with_next = True


def add_bullets(doc: Document, items: Iterable[str]) -> None:
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def add_landscape_section(doc: Document) -> None:
    section = doc.add_section()
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width
    section.top_margin = Cm(1.7)
    section.bottom_margin = Cm(1.7)
    section.left_margin = Cm(1.5)
    section.right_margin = Cm(1.5)


def add_portrait_section(doc: Document) -> None:
    section = doc.add_section()
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width, section.page_height = section.page_height, section.page_width
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.1)
    section.right_margin = Cm(2.1)


def setup_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_height = Cm(29.7)
    section.page_width = Cm(21.0)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.1)
    section.right_margin = Cm(2.1)

    normal = doc.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.2
    for name, size, color in (
        ("Title", 24, "000000"),
        ("Heading 1", 16, "000000"),
        ("Heading 2", 13, "000000"),
    ):
        style = doc.styles[name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)

    for section in doc.sections:
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        footer.add_run("ZS601 3DGS v3 · 证据驱动实验记录")


def build(registry: dict[str, Any], output: Path) -> None:
    doc = Document()
    setup_document(doc)

    title = doc.add_heading("ZS601 3DGS 实验总记录", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle = doc.add_paragraph("标准化数据预处理、验证协议与实验登记")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("生成时间：").bold = True
    p.add_run(datetime.now(timezone.utc).astimezone().isoformat(timespec="seconds"))
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("证据规则：").bold = True
    p.add_run(registry.get("evidence_policy", "缺少证据的字段标记为待核验。"))
    doc.add_page_break()

    add_heading(doc, "1. 项目输入数据集")
    doc.add_paragraph(
        "原始数据来自 L2 PRO 与 LCC 流程，包括透视 RGB、mask、相机标定和已经在上游去除动态目标的彩色 LiDAR 点云。"
        "Drive 逻辑路径为 MyDrive/LCCDataset/zs601_output/<scene_name>/。原始数据只读，派生数据写入 processed_v3。"
    )
    add_table(doc, ["数据角色", "要求", "身份记录"], [
        ["RGB", "名称唯一、尺寸与相机模型一致", "数量、格式、哈希"],
        ["mask", "与 RGB 一一对应，显式记录有效语义", "数量、编码、哈希"],
        ["相机", "COLMAP world-to-camera；C=-R^Tt", "模型、内外参、划分"],
        ["LiDAR", "已去动态目标；至少 x/y/z，通常含 RGB", "来源、字段、单位、SHA-256"],
    ])

    add_heading(doc, "2. SfM 数据预处理")
    add_bullets(doc, [
        "读取图像与内外参，统一为 COLMAP world-to-camera，并计算相机中心。",
        "按名称关联 RGB、mask 和姿态；检查重复、缺失和分辨率不一致。",
        "核对 cameras/images/points3D 的 bin 与 txt，记录训练实际读取格式。",
        "固化 train/val/test；val10 必须包含 10 个有效且不重复的相机。",
    ])

    add_heading(doc, "3. 去动态目标 LiDAR 点云处理")
    doc.add_paragraph(
        "本流程接收已经去除动态目标的点云，不再次删除动态目标。仅剔除 NaN、Inf 或无法解析点；额外离群点过滤默认关闭。"
        "使用 KNN/PCA 估计法向，保存曲率、邻居数、法向可靠度和 normal_valid。"
    )

    add_heading(doc, "4. 法向定向与逐相机法向图")
    doc.add_paragraph(
        "只使用训练相机筛选可见候选，为每点选择最近的最多 8 个相机。以距离、观察角和可见性加权点到相机方向；"
        "加权点积为负时翻转 PCA 法向。低置信度点可执行邻域符号传播，仍不确定者标记 unresolved。"
    )
    doc.add_paragraph(
        "定向后的世界法向通过 R 变换到相机坐标，与 z-depth 共用 z-buffer。浮点 NPY 保存 [-1,1]，PNG 使用 (n+1)/2 映射；"
        "独立 normal_valid 表示有效像素。报告区分未覆盖、mask 排除、低置信度和有效区近黑像素，避免把稀疏无效区误判为方向错误。"
    )

    add_heading(doc, "5. 深度生成与三维回投验证")
    doc.add_paragraph(
        "深度定义为相机坐标 z-depth。深度与法向使用同一次投影和 z-buffer。有效深度反投影到三维后，与其源 LiDAR 点比较，"
        "报告均值、中位数、P90、P95 和最大误差；阈值按场景单位和输入分辨率配置。"
    )

    add_heading(doc, "6. 实验背景与统一协议")
    protocol = registry.get("protocol", {})
    add_table(doc, ["项目", "协议值"], [
        ["正式迭代", protocol.get("target_iterations")],
        ["验证间隔", protocol.get("validation_interval")],
        ["预期验证时间点", protocol.get("expected_validation_timepoints")],
        ["checkpoint", protocol.get("expected_checkpoints")],
        ["固定验证相机", protocol.get("validation_camera_count")],
        ["正式 GPU", protocol.get("formal_gpu")],
        ["冒烟步数", protocol.get("smoke_iterations")],
    ])

    add_heading(doc, "7. 各组原计划与真实运行状态")
    add_landscape_section(doc)
    experiments = registry.get("experiments", [])
    add_table(doc, ["组", "计划目标", "计划功能", "状态", "运行 ID", "commit", "数据哈希", "证据"], [
        [
            exp.get("group"), exp.get("purpose"), exp.get("planned_features"),
            exp.get("status"), evidenced(exp,"run_id"), evidenced(exp,"git_commit"),
            evidenced(exp,"dataset_manifest_sha256"), exp.get("evidence"),
        ]
        for exp in experiments
    ])
    doc.add_paragraph("说明：状态、commit、数据哈希、Drive 路径和指标只有在 evidence 中存在可回查证据时才可改为已核验。")
    doc.add_paragraph(
        "original 组目前仍需与原始 train_mask.py 做 200 步逐项等价性验收；在该验收通过前，不能把新增 runner 中全部开关关闭直接写成原始 3DGS 等价。"
    )

    add_heading(doc, "8. 过程实验分析")
    doc.add_paragraph(
        "正式分析按 iteration 0、随后每 5000 步和最终步读取 val_metrics.csv、loss_log.csv、geometry_metrics.csv、training_progress.csv 及同轮次 RGB、1σ 椭球、深度和法向图。"
        "只有文件存在、迭代连续且运行身份一致时才写入下表。"
    )
    process_rows = []
    for exp in experiments:
        progress = exp.get("progress", [])
        if not progress:
            process_rows.append([exp.get("group"), None, None, None, None, "待核验"])
        else:
            for item in progress:
                process_rows.append([
                    exp.get("group"), item.get("iteration"), item.get("loss"),
                    item.get("gaussian_count"), item.get("geometry_metrics"), item.get("evidence"),
                ])
    add_table(doc, ["组", "iteration", "loss", "高斯数量", "几何指标", "证据"], process_rows)

    add_heading(doc, "9. 最终量化对比")
    doc.add_paragraph(
        "150000 步完成后，对全部显式 test 相机计算逐相机指标，保存 test_metrics_per_camera.csv；按 masked PSNR 升序保存最差 10 个相机的 RGB、1σ 椭球、深度和法向，并生成 CSV、Markdown 和 LaTeX 结果表。"
    )
    metric_keys = sorted({key for exp in experiments for key in exp.get("metrics", {})})
    if metric_keys:
        add_table(doc, ["组", *metric_keys, "证据状态"], [
            [exp.get("group"), *[
                exp.get("metrics", {}).get(k) if exp.get("evidence") else None
                for k in metric_keys],
             "已登记" if exp.get("evidence") else "待核验"]
            for exp in experiments
        ])
    else:
        doc.add_paragraph("当前 registry 未登记有证据支持的最终指标。所有最终量化结果均待核验。")

    add_portrait_section(doc)
    add_heading(doc, "10. 下一步启发")
    add_bullets(doc, [
        "先完成数据 manifest、投影叠加和深度反投影验证，再进入正式训练。",
        "用 iteration 0 的法向有效覆盖、近黑比例和邻域一致性验证相机辅助定向。",
        "每个新实验只改变登记中的控制变量，并预先写明成功标准和停止条件。",
        "正式运行完成后先回读 Drive 关键产物，再更新状态、指标和结论。",
    ])

    add_heading(doc, "附录 A：文档与证据来源")
    add_bullets(doc, [
        "docs/DATASET_INPUT.md",
        "docs/DATA_PREPROCESSING.md",
        "docs/VAL_DIAGNOSTICS.md",
        "experiments/experiment_registry.json",
        "仓库根目录 prompt/registers/experiment-state.md（更新 registry 前核验）",
        "仓库根目录 experiments/INDEX.md 与对应 summary（更新 registry 前核验）",
    ])

    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--registry", type=Path, default=DEFAULT_REGISTRY)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()
    with args.registry.open("r", encoding="utf-8") as handle:
        registry = json.load(handle)
    build(registry, args.output)
    print(args.output.resolve())


if __name__ == "__main__":
    main()
